from __future__ import annotations

import argparse
import html
import json
import re
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
MANIFEST = ROOT / "data" / "corpus" / "manifest.json"
DEFAULT_OUTPUT = ROOT / "data" / "corpus" / "generated"
FORMATS = {"md", "txt", "html", "docx", "pdf"}


def load_manifest(path: Path = MANIFEST) -> list[dict]:
    records = json.loads(path.read_text(encoding="utf-8"))
    seen: set[str] = set()
    for record in records:
        slug = record["id"]
        if not re.fullmatch(r"[a-z0-9-]+", slug) or slug in seen:
            raise ValueError(f"Invalid or duplicate document id: {slug}")
        if record["format"] not in FORMATS or not record["sections"] or len(record["qa"]) != 3:
            raise ValueError(f"Incomplete corpus record: {slug}")
        seen.add(slug)
    return records


def generate_corpus(output: Path = DEFAULT_OUTPUT) -> list[Path]:
    output.mkdir(parents=True, exist_ok=True)
    generated: list[Path] = []
    for record in load_manifest():
        path = output / f"{record['id']}.{record['format']}"
        title = record["title"]
        sections = record["sections"]
        if record["format"] == "md":
            content = f"# {title}\n\n" + "\n\n".join(f"## {heading}\n\n{body}" for heading, body in sections) + "\n"
            path.write_text(content, encoding="utf-8")
        elif record["format"] == "txt":
            content = title + "\n\n" + "\n\n".join(f"{heading}\n{body}" for heading, body in sections) + "\n"
            path.write_text(content, encoding="utf-8")
        elif record["format"] == "html":
            content = "<!doctype html><html lang=\"zh-CN\"><meta charset=\"utf-8\"><body>"
            content += f"<h1>{html.escape(title)}</h1>"
            content += "".join(f"<h2>{html.escape(heading)}</h2><p>{html.escape(body)}</p>" for heading, body in sections)
            path.write_text(content + "</body></html>\n", encoding="utf-8")
        elif record["format"] == "docx":
            from docx import Document

            document = Document()
            document.add_heading(title, level=1)
            for heading, body in sections:
                document.add_heading(heading, level=2)
                document.add_paragraph(body)
            document.save(path)
        else:
            from reportlab.lib.pagesizes import A4
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.cidfonts import UnicodeCIDFont
            from reportlab.pdfgen import canvas

            pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
            pdf = canvas.Canvas(str(path), pagesize=A4)
            pdf.setTitle(title)
            pdf.setFont("STSong-Light", 16)
            pdf.drawString(56, 790, title)
            y = 750
            for heading, body in sections:
                pdf.setFont("STSong-Light", 13)
                pdf.drawString(56, y, heading)
                y -= 27
                pdf.setFont("STSong-Light", 11)
                for start in range(0, len(body), 40):
                    pdf.drawString(56, y, body[start:start + 40])
                    y -= 20
                y -= 15
            pdf.save()
        generated.append(path)
    return generated


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Generate fictional enterprise policy files")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    paths = generate_corpus(args.output)
    print(f"Generated {len(paths)} fictional policies in {args.output}")
