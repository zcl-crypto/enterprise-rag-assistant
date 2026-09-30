from pathlib import Path

from docx import Document as WordDocument
from reportlab.pdfgen import canvas

from rag_app.documents import parse_file


def test_docx_and_html_parse_to_common_document(tmp_path: Path) -> None:
    word_path = tmp_path / "access.docx"
    word = WordDocument()
    word.add_heading("账号管理", level=1)
    word.add_paragraph("离职当天停用账号。")
    word.save(word_path)

    html_path = tmp_path / "budget.html"
    html_path.write_text("<html><body><h1>预算规则</h1><p>单笔超过一万元需审批。</p></body></html>", encoding="utf-8")

    for path, expected in ((word_path, "离职当天"), (html_path, "一万元")):
        parsed = parse_file(path)
        assert expected in "\n".join(block.text for block in parsed.blocks)
        assert parsed.source_format == path.suffix[1:]


def test_text_pdf_has_page_provenance(tmp_path: Path) -> None:
    pdf_path = tmp_path / "policy.pdf"
    pdf = canvas.Canvas(str(pdf_path))
    pdf.drawString(72, 750, "Travel policy: submit expenses within 30 days.")
    pdf.save()

    parsed = parse_file(pdf_path)

    assert "Travel policy" in "\n".join(block.text for block in parsed.blocks)
    assert any(block.page_number == 1 for block in parsed.blocks)
